from __future__ import annotations

from collections import defaultdict
from datetime import UTC, datetime
from html import escape
from io import BytesIO
import math
import re
from typing import Any

from sqlalchemy import desc, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.entities import (
    AssetModel,
    DataSource,
    EnvironmentObservation,
    Inspection,
    Maintenance,
    ModelRegistry,
    Observation,
    Prediction,
    Sensor,
)
from app.services.ml_predictor import condition_to_rating
from app.services.twin_service import build_twin, get_asset_row


GOVERNMENT_MARKERS = (
    "GOVERNMENT",
    "CENTRAL WATER COMMISSION",
    "CWC",
    "NWDP",
    "NWIC",
    "INDIA-WRIS",
    "WRIS",
    "MORTH",
    "MINISTRY OF ROAD TRANSPORT",
    "NHAI",
    "WATER RESOURCES",
    "INDIAN RAILWAYS",
    "RDSO",
    "AIRPORTS AUTHORITY",
    "AAI",
)


def _as_dict(value: Any) -> dict[str, Any]:
    return value if isinstance(value, dict) else {}


def _num(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def _iso(value: Any) -> str | None:
    if value is None:
        return None
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)


def _rounded(value: Any, digits: int = 2) -> float | None:
    number = _num(value)
    return round(number, digits) if number is not None else None


def _health_category(score: float | None) -> str:
    if score is None:
        return "NOT_AVAILABLE"
    if score >= 85:
        return "EXCELLENT"
    if score >= 70:
        return "GOOD"
    if score >= 50:
        return "WARNING"
    return "CRITICAL"


def _risk_level(score: float | None, supplied: Any = None) -> str:
    supplied_text = str(supplied or "").strip().upper()
    if supplied_text in {"LOW", "MEDIUM", "HIGH", "CRITICAL"}:
        return supplied_text
    if score is None:
        return "NOT_AVAILABLE"
    if score >= 85:
        return "CRITICAL"
    if score >= 70:
        return "HIGH"
    if score >= 40:
        return "MEDIUM"
    return "LOW"


def _percentile(values: list[float], q: float) -> float | None:
    clean = sorted(v for v in values if math.isfinite(v))
    if not clean:
        return None
    if len(clean) == 1:
        return clean[0]
    position = (len(clean) - 1) * q
    lo = int(math.floor(position))
    hi = int(math.ceil(position))
    if lo == hi:
        return clean[lo]
    weight = position - lo
    return clean[lo] * (1 - weight) + clean[hi] * weight


def _trend(points: list[dict[str, Any]]) -> str:
    values = [_num(point.get("value")) for point in points]
    usable = [value for value in values if value is not None]
    if len(usable) < 2:
        return "INSUFFICIENT_HISTORY"
    first = usable[0]
    last = usable[-1]
    span = max(abs(first), abs(last), 1.0)
    delta = last - first
    if abs(delta) <= 0.02 * span:
        return "STABLE"
    return "INCREASING" if delta > 0 else "DECREASING"


def _range_text(low: float | None, high: float | None, unit: str | None) -> str | None:
    if low is None or high is None:
        return None
    suffix = f" {unit}" if unit else ""
    return f"{low:.3g}–{high:.3g}{suffix}"


def _difference_from_range(
    value: float | None,
    low: float | None,
    high: float | None,
    unit: str | None,
) -> str | None:
    if value is None or low is None or high is None:
        return None
    suffix = f" {unit}" if unit else ""
    if value < low:
        return f"{low - value:.3g}{suffix} below historical range"
    if value > high:
        return f"{value - high:.3g}{suffix} above historical range"
    return "Within historical range"


def _is_government_source(source: DataSource | None) -> bool:
    if source is None or not bool(getattr(source, "is_authoritative", False)):
        return False
    text = " | ".join(
        str(getattr(source, field, "") or "")
        for field in ("code", "name", "organisation", "source_type")
    ).upper()
    return any(marker in text for marker in GOVERNMENT_MARKERS)


def _prediction_series(rows: list[Prediction], kind: str) -> list[dict[str, Any]]:
    output: list[dict[str, Any]] = []
    for row in rows:
        target = str(row.target or "").lower()
        matches = (
            (kind == "health" and "health" in target)
            or (kind == "risk" and "risk" in target)
            or (kind == "rul" and ("rul" in target or "remaining" in target))
        )
        if not matches:
            continue
        value = _num(row.value)
        if value is None:
            continue
        output.append(
            {
                "timestamp": _iso(row.prediction_time),
                "value": round(value, 3),
                "lower_bound": _rounded(row.lower_bound, 3),
                "upper_bound": _rounded(row.upper_bound, 3),
                "confidence": _rounded(row.confidence_score, 3),
                "model_version": row.model_version,
                "status": row.status,
            }
        )
    return output[-40:]


def _supporting_row(
    *,
    name: str,
    current_value: float | str | None,
    unit: str | None,
    history: list[dict[str, Any]],
    status: str | None,
    updated_at: Any,
    source: str | None,
    used_by_model: bool,
) -> dict[str, Any]:
    numeric_history = [
        number
        for point in history
        if (number := _num(point.get("value"))) is not None
    ]
    low = _percentile(numeric_history, 0.05) if len(numeric_history) >= 5 else None
    high = _percentile(numeric_history, 0.95) if len(numeric_history) >= 5 else None
    return {
        "name": name,
        "current_value": current_value,
        "unit": unit,
        "expected_range": _range_text(low, high, unit),
        "reference_type": (
            "HISTORICAL_EMPIRICAL_P05_P95"
            if low is not None and high is not None
            else "NOT_AVAILABLE"
        ),
        "difference_from_expected": _difference_from_range(
            _num(current_value),
            low,
            high,
            unit,
        ),
        "status": status,
        "trend": _trend(history),
        "last_updated": _iso(updated_at),
        "source": source,
        "used_by_model": used_by_model,
        "history": history[-60:],
    }


def _model_factor_rows(
    *,
    asset: Any,
    dimensions: dict[str, Any],
    inspection: Inspection | None,
    ai: dict[str, Any],
) -> tuple[list[dict[str, Any]], int]:
    # Never infer a structural rating from asset.condition here. The full
    # structural model input is shown only when an eligible non-synthetic
    # inspection is actually available.
    condition_rating = (
        condition_to_rating(inspection.condition, inspection.score)
        if inspection is not None
        else None
    )
    age_years = (
        datetime.now(UTC).year - int(asset.built_year)
        if getattr(asset, "built_year", None)
        else None
    )

    method = str(ai.get("prediction_method") or "").lower()
    full_bridge = method == "calibrated_nbi_bridge_ml"
    sparse_bridge = "sparse_nbi_bridge_ml" in method

    factors = [
        {
            "factor": "Age",
            "current_value": age_years,
            "unit": "years",
            "expected_range": "0–200 model input domain",
            "source": "Canonical asset built year",
            "used_by_model": (full_bridge or sparse_bridge) and age_years is not None,
        },
        {
            "factor": "Structural condition rating",
            "current_value": condition_rating,
            "unit": "/9",
            "expected_range": "0–9 model input domain",
            "source": "Latest eligible non-synthetic inspection",
            "used_by_model": full_bridge and condition_rating is not None,
        },
        {
            "factor": "Material",
            "current_value": getattr(asset, "material", None),
            "unit": None,
            "expected_range": None,
            "source": "Canonical asset record",
            "used_by_model": (
                full_bridge or sparse_bridge
            ) and getattr(asset, "material", None) is not None,
        },
        {
            "factor": "Span count",
            "current_value": dimensions.get("span_count"),
            "unit": None,
            "expected_range": None,
            "source": "Active asset-model dimensions",
            "used_by_model": (
                full_bridge or sparse_bridge
            ) and dimensions.get("span_count") is not None,
        },
        {
            "factor": "Maximum / main span",
            "current_value": dimensions.get("main_span_m") or dimensions.get("max_span_m"),
            "unit": "m",
            "expected_range": None,
            "source": "Active asset-model dimensions",
            "used_by_model": (
                full_bridge or sparse_bridge
            ) and (
                dimensions.get("main_span_m") is not None
                or dimensions.get("max_span_m") is not None
            ),
        },
        {
            "factor": "Structure length",
            "current_value": dimensions.get("length_m"),
            "unit": "m",
            "expected_range": None,
            "source": "Active asset-model dimensions",
            "used_by_model": (
                full_bridge or sparse_bridge
            ) and dimensions.get("length_m") is not None,
        },
    ]

    for item in factors:
        item["difference_from_expected"] = None
        item["contribution"] = None
        item["impact"] = (
            "Used by the active model. Signed feature contribution is not stored in the current artifact."
            if item["used_by_model"]
            else "Available context; not used by the active structural model."
        )

    return factors, sum(1 for item in factors if item["used_by_model"])


def _recommendations(ai: dict[str, Any]) -> list[dict[str, Any]]:
    raw = ai.get("recommendations")
    if not isinstance(raw, list):
        return []

    current_risk = _num(ai.get("risk_score"))
    if current_risk is None:
        current_risk = _num(ai.get("operational_risk_score"))
    current_level = _risk_level(
        current_risk,
        ai.get("risk_level") or ai.get("operational_risk_level"),
    )
    priority = (
        current_level if current_level in {"LOW", "MEDIUM", "HIGH", "CRITICAL"} else "MEDIUM"
    )
    factors = [str(item) for item in ai.get("factors", []) if str(item).strip()]

    output: list[dict[str, Any]] = []
    for value in raw:
        recommendation = str(value).strip()
        if not recommendation:
            continue

        timeframe = None
        match = re.search(
            r"\bwithin\s+(\d+)\s+(day|days|week|weeks|month|months|year|years)\b",
            recommendation,
            re.I,
        )
        if match:
            timeframe = match.group(0)
        elif re.search(r"\bimmediate(ly)?\b|\burgent\b", recommendation, re.I):
            timeframe = "Immediate / urgent"

        output.append(
            {
                "recommendation": recommendation,
                "priority": priority,
                "priority_basis": "Current backend risk level",
                "reason": (
                    factors[0]
                    if factors
                    else "Current SIMRAS backend recommendation logic returned this action."
                ),
                "supporting_evidence": factors[:4],
                "expected_benefit": None,
                "suggested_timeframe": timeframe,
                "source": "Current SIMRAS backend recommendation engine",
            }
        )
    return output


async def build_asset_health_assessment(
    session: AsyncSession,
    asset_code: str,
) -> dict[str, Any]:
    row = await get_asset_row(session, asset_code)
    asset, longitude, latitude = row
    twin = await build_twin(session, asset_code)
    ai = _as_dict(twin.get("ai"))
    now = datetime.now(UTC)

    model = (
        await session.execute(
            select(AssetModel)
            .where(
                AssetModel.asset_id == asset.id,
                AssetModel.is_active.is_(True),
            )
            .order_by(desc(AssetModel.updated_at))
            .limit(1)
        )
    ).scalar_one_or_none()
    dimensions = _as_dict(getattr(model, "dimensions", None))

    inspection_rows = (
        await session.execute(
            select(Inspection)
            .where(Inspection.asset_id == asset.id)
            .order_by(desc(Inspection.inspection_date), desc(Inspection.id))
            .limit(50)
        )
    ).scalars().all()
    eligible_inspection = next(
        (
            item
            for item in inspection_rows
            if not bool(getattr(item, "is_synthetic", False))
        ),
        None,
    )

    maintenance_rows = (
        await session.execute(
            select(Maintenance)
            .where(Maintenance.asset_id == asset.id)
            .order_by(desc(Maintenance.maintenance_date), desc(Maintenance.id))
            .limit(50)
        )
    ).scalars().all()

    prediction_rows = (
        await session.execute(
            select(Prediction)
            .where(Prediction.asset_id == asset.id)
            .order_by(Prediction.prediction_time.asc(), Prediction.id.asc())
        )
    ).scalars().all()

    sensor_rows = (
        await session.execute(
            select(Sensor)
            .where(Sensor.asset_id == asset.id)
            .order_by(Sensor.observed_property, Sensor.sensor_code)
        )
    ).scalars().all()

    sensor_support: list[dict[str, Any]] = []
    for sensor in sensor_rows:
        observations = (
            await session.execute(
                select(Observation)
                .where(Observation.sensor_id == sensor.id)
                .order_by(desc(Observation.observed_at), desc(Observation.id))
                .limit(120)
            )
        ).scalars().all()
        observations = list(reversed(observations))
        history = [
            {
                "timestamp": _iso(obs.observed_at),
                "value": _rounded(obs.value, 4),
                "quality_flag": obs.quality_flag,
            }
            for obs in observations
        ]
        latest = observations[-1] if observations else None
        sensor_support.append(
            _supporting_row(
                name=sensor.observed_property,
                current_value=_rounded(latest.value, 4) if latest else None,
                unit=sensor.unit,
                history=history,
                status=latest.quality_flag if latest else sensor.status,
                updated_at=latest.observed_at if latest else sensor.updated_at,
                source=f"Sensor {sensor.sensor_code}",
                used_by_model=False,
            )
        )

    env_pairs = (
        await session.execute(
            select(EnvironmentObservation, DataSource)
            .join(DataSource, DataSource.id == EnvironmentObservation.source_id)
            .where(EnvironmentObservation.asset_id == asset.id)
            .order_by(
                EnvironmentObservation.variable,
                desc(EnvironmentObservation.observed_at),
            )
            .limit(2500)
        )
    ).all()

    env_grouped: dict[str, list[tuple[EnvironmentObservation, DataSource]]] = defaultdict(list)
    for observation, source in env_pairs:
        env_grouped[observation.variable].append((observation, source))

    environment_support: list[dict[str, Any]] = []
    for variable, pairs in sorted(env_grouped.items()):
        chronological = list(reversed(pairs[:120]))
        history = [
            {
                "timestamp": _iso(observation.observed_at),
                "value": _rounded(observation.value, 4),
                "quality_flag": observation.quality_flag,
                "source": source.name,
            }
            for observation, source in chronological
        ]
        latest_observation, latest_source = pairs[0]
        environment_support.append(
            _supporting_row(
                name=variable,
                current_value=_rounded(latest_observation.value, 4),
                unit=latest_observation.unit,
                history=history,
                status=latest_observation.quality_flag,
                updated_at=latest_observation.observed_at,
                source=latest_source.name,
                used_by_model=False,
            )
        )

    model_inputs, model_input_count = _model_factor_rows(
        asset=asset,
        dimensions=dimensions,
        inspection=eligible_inspection,
        ai=ai,
    )
    explanation_notes = [
        str(item)
        for item in ai.get("factors", [])
        if str(item).strip()
    ]

    health_score = _num(ai.get("health_score"))
    structural_risk = _num(ai.get("risk_score"))
    operational_risk = _num(ai.get("operational_risk_score"))
    displayed_risk = structural_risk if structural_risk is not None else operational_risk
    displayed_level = _risk_level(
        displayed_risk,
        ai.get("risk_level") or ai.get("operational_risk_level"),
    )
    confidence = _num(ai.get("confidence"))
    if confidence is None:
        confidence = _num(ai.get("operational_confidence"))

    health_history = _prediction_series(prediction_rows, "health")
    risk_history = _prediction_series(prediction_rows, "risk")
    rul_history = _prediction_series(prediction_rows, "rul")

    latest_prediction_time: Any = (
        prediction_rows[-1].prediction_time if prediction_rows else None
    )
    if latest_prediction_time is None:
        latest_prediction_time = ai.get("prediction_time") or twin.get("generated_at")

    model_version = ai.get("model_version")
    registry = None
    if model_version:
        registry = (
            await session.execute(
                select(ModelRegistry)
                .where(ModelRegistry.version == str(model_version))
                .order_by(desc(ModelRegistry.id))
                .limit(1)
            )
        ).scalar_one_or_none()

    source_ids = {
        int(source_id)
        for source_id in (
            [getattr(asset, "source_id", None)]
            + [item.source_id for item in inspection_rows]
            + [item.source_id for item in maintenance_rows]
            + [observation.source_id for observation, _ in env_pairs]
        )
        if source_id is not None
    }
    source_cache: dict[int, DataSource] = {}
    if source_ids:
        sources = (
            await session.execute(
                select(DataSource).where(DataSource.id.in_(source_ids))
            )
        ).scalars().all()
        source_cache = {source.id: source for source in sources}

    government_rows: list[dict[str, Any]] = []

    asset_source = source_cache.get(getattr(asset, "source_id", None))
    if _is_government_source(asset_source):
        government_rows.append(
            {
                "kind": "AUTHORITATIVE_ASSET_RECORD",
                "regulation_or_standard": None,
                "evidence_or_requirement": "Authoritative asset identity / registry record",
                "current_measured_value": asset.status,
                "applicable_threshold": None,
                "compliance_status": "NOT_ASSESSED_NO_STORED_THRESHOLD",
                "source_reference": f"{asset_source.code} · {asset_source.name}",
                "reference_url": asset_source.url,
                "observed_at": _iso(getattr(asset, "updated_at", None)),
            }
        )

    for inspection in inspection_rows[:10]:
        source = source_cache.get(inspection.source_id)
        if not _is_government_source(source):
            continue
        measured = inspection.condition
        if inspection.score is not None:
            measured = f"{inspection.condition} · score {inspection.score:g}"
        government_rows.append(
            {
                "kind": "AUTHORITATIVE_INSPECTION_RECORD",
                "regulation_or_standard": None,
                "evidence_or_requirement": f"{inspection.inspection_type} inspection record",
                "current_measured_value": measured,
                "applicable_threshold": None,
                "compliance_status": "NOT_ASSESSED_NO_STORED_THRESHOLD",
                "source_reference": f"{source.code} · {source.name}",
                "reference_url": source.url,
                "observed_at": _iso(inspection.inspection_date),
            }
        )

    seen_environment_variables: set[str] = set()
    for observation, source in env_pairs:
        if observation.variable in seen_environment_variables:
            continue
        if not _is_government_source(source):
            continue
        seen_environment_variables.add(observation.variable)
        government_rows.append(
            {
                "kind": "AUTHORITATIVE_OBSERVATION",
                "regulation_or_standard": None,
                "evidence_or_requirement": observation.variable.replace("_", " "),
                "current_measured_value": f"{observation.value:g} {observation.unit}".strip(),
                "applicable_threshold": None,
                "compliance_status": "NOT_ASSESSED_NO_STORED_THRESHOLD",
                "source_reference": f"{source.code} · {source.name}",
                "reference_url": source.url,
                "observed_at": _iso(observation.observed_at),
            }
        )

    unique_government_rows: list[dict[str, Any]] = []
    seen_government_rows: set[tuple[str, str, str]] = set()
    for item in government_rows:
        key = (
            str(item.get("kind")),
            str(item.get("evidence_or_requirement")),
            str(item.get("source_reference")),
        )
        if key in seen_government_rows:
            continue
        seen_government_rows.add(key)
        unique_government_rows.append(item)

    prediction_method = str(ai.get("prediction_method") or "")
    poor_condition_probability = (
        structural_risk
        if (
            prediction_method == "calibrated_nbi_bridge_ml"
            or "sparse_nbi_bridge_ml" in prediction_method
        )
        else None
    )

    report_time_seed = _iso(latest_prediction_time) or now.isoformat()
    timestamp_digits = re.sub(r"[^0-9]", "", report_time_seed)[:14]
    if not timestamp_digits:
        timestamp_digits = now.strftime("%Y%m%d%H%M%S")
    report_id = f"SIMRAS-{asset.asset_code}-{timestamp_digits}"

    supporting_data = sensor_support + environment_support
    evidence_record_count = (
        len([item for item in inspection_rows if not bool(item.is_synthetic)])
        + len([item for item in maintenance_rows if not bool(item.is_synthetic)])
        + sum(len(item["history"]) for item in supporting_data)
    )

    return {
        "schema_version": "simras-asset-health-assessment-v1",
        "report_id": report_id,
        "generated_at": now.isoformat(),
        "asset": {
            "id": asset.id,
            "asset_code": asset.asset_code,
            "name": asset.name,
            "asset_type": asset.asset_type,
            "subtype": asset.subtype,
            "district": asset.district,
            "owner": asset.owner,
            "status": asset.status,
            "identity_status": asset.identity_status,
            "built_year": asset.built_year,
            "material": asset.material,
            "longitude": _rounded(longitude, 7),
            "latitude": _rounded(latitude, 7),
        },
        "health": {
            "score": _rounded(health_score, 0),
            "raw_score": _rounded(health_score, 4),
            "category": _health_category(health_score),
            "lower_bound": _rounded(ai.get("health_lower_bound"), 0),
            "upper_bound": _rounded(ai.get("health_upper_bound"), 0),
            "factors": explanation_notes,
            "trend": _trend(health_history),
            "history": health_history,
            "available": health_score is not None,
        },
        "risk": {
            "score": _rounded(displayed_risk, 0),
            "raw_score": _rounded(displayed_risk, 4),
            "level": displayed_level,
            "confidence": _rounded(confidence, 3),
            "poor_condition_probability": _rounded(poor_condition_probability, 2),
            "failure_probability": None,
            "failure_probability_note": (
                "The active bridge model predicts poor-condition probability within its forecast horizon; "
                "it does not expose a separately validated structural failure probability."
                if poor_condition_probability is not None
                else "No separately validated structural failure-probability output is stored for this prediction."
            ),
            "factors": explanation_notes,
            "trend": _trend(risk_history),
            "history": risk_history,
            "available": displayed_risk is not None,
        },
        "rul": {
            "estimate": _rounded(ai.get("remaining_life_years"), 0),
            "raw_estimate": _rounded(ai.get("remaining_life_years"), 4),
            "unit": "years",
            "lower_bound": _rounded(ai.get("rul_lower_bound"), 0),
            "upper_bound": _rounded(ai.get("rul_upper_bound"), 0),
            "confidence": _rounded(ai.get("rul_confidence") or confidence, 3),
            "status": ai.get("rul_status") or ai.get("status"),
            "basis": ai.get("rul_basis"),
            "trend": _trend(rul_history),
            "history": rul_history,
            "available": _num(ai.get("remaining_life_years")) is not None,
        },
        "explainability": {
            "model_inputs": model_inputs,
            "model_explanation_notes": explanation_notes,
            "signed_contributions_available": False,
            "signed_contribution_note": (
                "The currently registered artifact does not store SHAP/signed feature contributions. "
                "This report exposes actual model inputs and backend factors without inventing attribution."
            ),
        },
        "government_evidence": {
            "records": unique_government_rows,
            "verified_record_count": len(unique_government_rows),
            "verified_standard_count": sum(
                1 for item in unique_government_rows if item.get("regulation_or_standard")
            ),
            "message": (
                None
                if unique_government_rows
                else "No verified government/regulatory evidence available for this prediction."
            ),
            "threshold_note": (
                "No regulatory threshold is inferred. Compliance is assessed only when a verified threshold/standard is explicitly stored."
            ),
        },
        "recommendations": _recommendations(ai),
        "supporting_data": supporting_data,
        "inspection_history": [
            {
                "date": _iso(item.inspection_date),
                "type": item.inspection_type,
                "condition": item.condition,
                "score": _rounded(item.score, 2),
                "quality_flag": item.quality_flag,
                "is_synthetic": bool(item.is_synthetic),
                "source": (
                    source_cache[item.source_id].name
                    if item.source_id in source_cache
                    else None
                ),
            }
            for item in inspection_rows
        ],
        "maintenance_history": [
            {
                "date": _iso(item.maintenance_date),
                "action": item.action,
                "status": item.status,
                "next_due": _iso(item.next_due),
                "is_synthetic": bool(item.is_synthetic),
                "source": (
                    source_cache[item.source_id].name
                    if item.source_id in source_cache
                    else None
                ),
            }
            for item in maintenance_rows
        ],
        "transparency": {
            "prediction": displayed_level if displayed_risk is not None else "NOT_AVAILABLE",
            "confidence": _rounded(confidence, 3),
            "model_name": getattr(registry, "model_name", None) or ai.get("model_name"),
            "model_version": model_version or getattr(registry, "version", None),
            "feature_version": ai.get("feature_version") or getattr(registry, "feature_version", None),
            "model_stage": ai.get("status") or getattr(registry, "stage", None),
            "prediction_method": prediction_method or None,
            "prediction_generated_at": _iso(latest_prediction_time),
            "data_points_used": model_input_count,
            "model_input_count": model_input_count,
            "evidence_record_count": evidence_record_count,
            "model_validated": bool(ai.get("model_validated")),
            "training_scope": ai.get("training_scope"),
        },
        "quality_boundaries": {
            "no_fake_values": True,
            "missing_values_remain_null": True,
            "government_thresholds_not_invented": True,
            "signed_feature_contributions_not_invented": True,
            "failure_probability_not_relabelled_from_poor_condition_probability": True,
        },
    }


def _chart_png(
    title: str,
    history: list[dict[str, Any]],
    ylabel: str,
) -> bytes | None:
    usable = [
        (point.get("timestamp"), _num(point.get("value")))
        for point in history
        if point.get("timestamp") and _num(point.get("value")) is not None
    ]
    if len(usable) < 2:
        return None

    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    x = list(range(len(usable)))
    y = [float(value) for _, value in usable if value is not None]
    labels = [str(timestamp)[:10] for timestamp, _ in usable]

    figure = plt.figure(figsize=(7.3, 2.6))
    axis = figure.add_subplot(111)
    axis.plot(x, y, marker="o", linewidth=1.6)
    axis.set_title(title)
    axis.set_ylabel(ylabel)
    axis.grid(True, alpha=0.25)

    if len(x) <= 12:
        ticks = x
    else:
        step = max(1, len(x) // 8)
        ticks = x[::step]

    axis.set_xticks(ticks)
    axis.set_xticklabels(
        [labels[index] for index in ticks],
        rotation=35,
        ha="right",
        fontsize=7,
    )
    figure.tight_layout()

    output = BytesIO()
    figure.savefig(output, format="png", dpi=150)
    plt.close(figure)
    return output.getvalue()


def assessment_pdf_bytes(report: dict[str, Any]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        Image,
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=14 * mm,
        leftMargin=14 * mm,
        topMargin=14 * mm,
        bottomMargin=14 * mm,
        title=f"{report['asset']['name']} - SIMRAS Asset Health Assessment",
    )
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="SIMRASTitle",
            parent=styles["Title"],
            alignment=TA_CENTER,
            fontSize=18,
            spaceAfter=8,
        )
    )
    styles.add(
        ParagraphStyle(
            name="SIMRASSmall",
            parent=styles["BodyText"],
            fontSize=8.3,
            leading=10.5,
        )
    )
    styles.add(
        ParagraphStyle(
            name="SIMRASSection",
            parent=styles["Heading2"],
            fontSize=13,
            spaceBefore=8,
            spaceAfter=6,
        )
    )

    story: list[Any] = []
    asset = report["asset"]
    story.append(
        Paragraph(
            "SIMRAS AI-Powered Asset Health Assessment Report",
            styles["SIMRASTitle"],
        )
    )

    metadata = Table(
        [
            ["Report ID", report["report_id"]],
            ["Asset", f"{asset['asset_code']} · {asset['name']}"],
            ["Generated", report["generated_at"]],
            ["Current asset status", asset.get("status") or "N/A"],
        ],
        colWidths=[42 * mm, 132 * mm],
    )
    metadata.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#e8f4f8")),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9bb8c6")),
                ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("PADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    story.extend([metadata, Spacer(1, 8)])

    health = report["health"]
    risk = report["risk"]
    rul = report["rul"]
    metrics = [
        ["Metric", "Value", "Status / interval"],
        [
            "Health score",
            f"{health['score']}/100" if health.get("score") is not None else "N/A",
            health.get("category") or "N/A",
        ],
        [
            "Risk",
            f"{risk['score']}/100" if risk.get("score") is not None else "N/A",
            risk.get("level") or "N/A",
        ],
        [
            "RUL",
            (
                f"{rul['estimate']} {rul['unit']}"
                if rul.get("estimate") is not None
                else "N/A"
            ),
            (
                f"{rul['lower_bound']}–{rul['upper_bound']} {rul['unit']}"
                if rul.get("lower_bound") is not None and rul.get("upper_bound") is not None
                else "N/A"
            ),
        ],
        [
            "Confidence",
            (
                f"{round(report['transparency']['confidence'] * 100)}%"
                if report["transparency"].get("confidence") is not None
                else "N/A"
            ),
            report["transparency"].get("model_stage") or "N/A",
        ],
    ]
    metrics_table = Table(metrics, colWidths=[52 * mm, 40 * mm, 82 * mm])
    metrics_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0d5268")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#aac3cd")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                ("PADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    story.append(metrics_table)

    for title, key, unit in (
        ("Health score trend", "health", "Health /100"),
        ("Risk trend", "risk", "Risk /100"),
        ("RUL trend", "rul", "Years"),
    ):
        chart = _chart_png(title, report[key].get("history", []), unit)
        if chart:
            story.extend(
                [
                    Spacer(1, 8),
                    Image(BytesIO(chart), width=174 * mm, height=62 * mm),
                ]
            )

    story.append(
        Paragraph(
            "Why did the AI make this prediction?",
            styles["SIMRASSection"],
        )
    )
    rows = [["Factor", "Current value", "Expected / reference", "Impact"]]
    for item in report["explainability"]["model_inputs"]:
        current = item.get("current_value")
        if current is not None and item.get("unit"):
            current = f"{current} {item['unit']}"
        rows.append(
            [
                item.get("factor") or "N/A",
                str(current if current is not None else "N/A"),
                item.get("expected_range") or "N/A",
                item.get("impact") or "N/A",
            ]
        )
    table = Table(
        rows,
        colWidths=[38 * mm, 34 * mm, 45 * mm, 57 * mm],
        repeatRows=1,
    )
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dceef3")),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#a6bdc7")),
                ("FONTSIZE", (0, 0), (-1, -1), 7.3),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("PADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story.append(table)

    for note in report["explainability"].get("model_explanation_notes", []):
        story.append(
            Paragraph(
                "• " + escape(str(note)),
                styles["SIMRASSmall"],
            )
        )
    story.append(
        Paragraph(
            escape(report["explainability"]["signed_contribution_note"]),
            styles["SIMRASSmall"],
        )
    )

    story.append(
        Paragraph(
            "Government Evidence & Standards",
            styles["SIMRASSection"],
        )
    )
    government = report["government_evidence"]
    if not government["records"]:
        story.append(
            Paragraph(
                "No verified government/regulatory evidence available for this prediction.",
                styles["BodyText"],
            )
        )
    else:
        rows = [
            [
                "Evidence / requirement",
                "Measured value",
                "Threshold",
                "Compliance",
                "Source",
            ]
        ]
        for item in government["records"]:
            rows.append(
                [
                    item.get("evidence_or_requirement") or "N/A",
                    item.get("current_measured_value") or "N/A",
                    item.get("applicable_threshold") or "N/A",
                    item.get("compliance_status") or "N/A",
                    item.get("source_reference") or "N/A",
                ]
            )
        table = Table(
            rows,
            colWidths=[42 * mm, 32 * mm, 28 * mm, 37 * mm, 35 * mm],
            repeatRows=1,
        )
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dceef3")),
                    ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#a6bdc7")),
                    ("FONTSIZE", (0, 0), (-1, -1), 7),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("PADDING", (0, 0), (-1, -1), 3),
                ]
            )
        )
        story.append(table)
        story.append(
            Paragraph(
                escape(government["threshold_note"]),
                styles["SIMRASSmall"],
            )
        )

    story.append(Paragraph("Recommendations", styles["SIMRASSection"]))
    if not report["recommendations"]:
        story.append(
            Paragraph(
                "No recommendation is available from the current backend recommendation logic.",
                styles["BodyText"],
            )
        )
    else:
        for item in report["recommendations"]:
            story.append(
                Paragraph(
                    f"<b>{escape(item['priority'])}</b> · {escape(item['recommendation'])}",
                    styles["BodyText"],
                )
            )
            story.append(
                Paragraph(
                    "Reason: " + escape(item["reason"]),
                    styles["SIMRASSmall"],
                )
            )
            story.append(
                Paragraph(
                    "Supporting evidence: "
                    + escape(" | ".join(item.get("supporting_evidence") or []) or "N/A"),
                    styles["SIMRASSmall"],
                )
            )
            story.append(
                Paragraph(
                    "Expected benefit: "
                    + escape(
                        item.get("expected_benefit")
                        or "Not quantified by current recommendation logic"
                    ),
                    styles["SIMRASSmall"],
                )
            )
            story.append(
                Paragraph(
                    "Suggested timeframe: "
                    + escape(
                        item.get("suggested_timeframe")
                        or "Not specified by current recommendation logic"
                    ),
                    styles["SIMRASSmall"],
                )
            )
            story.append(Spacer(1, 4))

    story.append(PageBreak())
    story.append(Paragraph("Supporting Data", styles["SIMRASSection"]))

    support_rows = [
        [
            "Input",
            "Current",
            "Expected / historical range",
            "Status",
            "Trend",
            "Updated",
            "Source",
        ]
    ]
    for item in report["supporting_data"]:
        current = item.get("current_value")
        if current is not None and item.get("unit"):
            current = f"{current} {item['unit']}"
        support_rows.append(
            [
                item.get("name") or "N/A",
                str(current if current is not None else "N/A"),
                item.get("expected_range") or "N/A",
                item.get("status") or "N/A",
                item.get("trend") or "N/A",
                item.get("last_updated") or "N/A",
                item.get("source") or "N/A",
            ]
        )

    if len(support_rows) > 1:
        table = Table(
            support_rows,
            colWidths=[
                29 * mm,
                24 * mm,
                34 * mm,
                24 * mm,
                22 * mm,
                34 * mm,
                28 * mm,
            ],
            repeatRows=1,
        )
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dceef3")),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#a6bdc7")),
                    ("FONTSIZE", (0, 0), (-1, -1), 6.3),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("PADDING", (0, 0), (-1, -1), 2.5),
                ]
            )
        )
        story.append(table)
    else:
        story.append(
            Paragraph(
                "No sensor/environment supporting-data rows are available.",
                styles["BodyText"],
            )
        )

    for item in report["supporting_data"][:8]:
        chart = _chart_png(
            item["name"],
            item.get("history", []),
            item.get("unit") or "Value",
        )
        if chart:
            story.extend(
                [
                    Spacer(1, 7),
                    Image(BytesIO(chart), width=174 * mm, height=62 * mm),
                ]
            )

    story.append(
        Paragraph(
            "Prediction Transparency",
            styles["SIMRASSection"],
        )
    )
    transparency = report["transparency"]
    transparency_rows = [
        ["Prediction", transparency.get("prediction") or "N/A"],
        [
            "Confidence",
            (
                f"{round(transparency['confidence'] * 100)}%"
                if transparency.get("confidence") is not None
                else "N/A"
            ),
        ],
        [
            "Model",
            f"{transparency.get('model_name') or 'N/A'} · {transparency.get('model_version') or 'N/A'}",
        ],
        ["Prediction method", transparency.get("prediction_method") or "N/A"],
        ["Prediction generated", transparency.get("prediction_generated_at") or "N/A"],
        ["Data points used", str(transparency.get("data_points_used") or 0)],
        ["Evidence records available", str(transparency.get("evidence_record_count") or 0)],
        ["Model stage", transparency.get("model_stage") or "N/A"],
    ]
    table = Table(transparency_rows, colWidths=[55 * mm, 119 * mm])
    table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#a6bdc7")),
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#edf6f8")),
                ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("PADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story.append(table)
    story.extend(
        [
            Spacer(1, 8),
            Paragraph(
                "Traceability boundary: populated values come from the current database, canonical Twin/ML output, stored histories, and authoritative source metadata. Missing regulatory thresholds, structural failure probabilities, and signed XAI contributions are not fabricated.",
                styles["SIMRASSmall"],
            ),
        ]
    )

    document.build(story)
    return output.getvalue()


def assessment_docx_bytes(report: dict[str, Any]) -> bytes:
    from docx import Document
    from docx.shared import Inches, Pt

    document = Document()
    document.add_heading(
        "SIMRAS AI-Powered Asset Health Assessment Report",
        0,
    )
    asset = report["asset"]

    metadata = document.add_table(rows=0, cols=2)
    for name, value in (
        ("Report ID", report["report_id"]),
        ("Asset", f"{asset['asset_code']} · {asset['name']}"),
        ("Generated", report["generated_at"]),
        ("Current asset status", asset.get("status") or "N/A"),
    ):
        cells = metadata.add_row().cells
        cells[0].text = name
        cells[1].text = str(value)

    health = report["health"]
    document.add_heading("Health Score", level=1)
    paragraph = document.add_paragraph()
    paragraph.add_run(
        f"{health['score']}/100" if health.get("score") is not None else "N/A"
    ).bold = True
    paragraph.add_run(f" · {health.get('category') or 'N/A'}")
    if health.get("lower_bound") is not None and health.get("upper_bound") is not None:
        document.add_paragraph(
            f"Prediction interval: {health['lower_bound']}–{health['upper_bound']}/100"
        )
    for factor in health.get("factors", []):
        document.add_paragraph(str(factor), style="List Bullet")

    risk = report["risk"]
    document.add_heading("Risk Assessment", level=1)
    document.add_paragraph(
        f"Risk: {risk['score']}/100 · {risk['level']}"
        if risk.get("score") is not None
        else f"Risk: N/A · {risk['level']}"
    )
    document.add_paragraph(
        f"Risk confidence: {round(risk['confidence'] * 100)}%"
        if risk.get("confidence") is not None
        else "Risk confidence: N/A"
    )
    document.add_paragraph(
        "Poor-condition probability: "
        + (
            f"{risk['poor_condition_probability']}%"
            if risk.get("poor_condition_probability") is not None
            else "N/A"
        )
    )
    document.add_paragraph(
        "Structural failure probability: N/A"
        if risk.get("failure_probability") is None
        else f"Structural failure probability: {risk['failure_probability']}%"
    )
    document.add_paragraph(risk.get("failure_probability_note") or "")

    rul = report["rul"]
    document.add_heading("Remaining Useful Life (RUL)", level=1)
    if rul.get("estimate") is None:
        document.add_paragraph("RUL: N/A")
    else:
        document.add_paragraph(f"RUL: {rul['estimate']} {rul['unit']}")
        if rul.get("lower_bound") is not None and rul.get("upper_bound") is not None:
            document.add_paragraph(
                f"Confidence range: {rul['lower_bound']}–{rul['upper_bound']} {rul['unit']}"
            )
    if rul.get("basis"):
        document.add_paragraph(str(rul["basis"]))

    for title, key, unit in (
        ("Health score trend", "health", "Health /100"),
        ("Risk trend", "risk", "Risk /100"),
        ("RUL trend", "rul", "Years"),
    ):
        chart = _chart_png(
            title,
            report[key].get("history", []),
            unit,
        )
        if chart:
            document.add_heading(title, level=2)
            document.add_picture(BytesIO(chart), width=Inches(6.4))

    document.add_heading(
        "Why did the AI make this prediction?",
        level=1,
    )
    table = document.add_table(rows=1, cols=4)
    for index, value in enumerate(
        ["Factor", "Current value", "Expected / reference", "Impact"]
    ):
        table.rows[0].cells[index].text = value

    for item in report["explainability"]["model_inputs"]:
        cells = table.add_row().cells
        current = item.get("current_value")
        if current is not None and item.get("unit"):
            current = f"{current} {item['unit']}"
        cells[0].text = str(item.get("factor") or "N/A")
        cells[1].text = str(current if current is not None else "N/A")
        cells[2].text = str(item.get("expected_range") or "N/A")
        cells[3].text = str(item.get("impact") or "N/A")

    for note in report["explainability"].get("model_explanation_notes", []):
        document.add_paragraph(str(note), style="List Bullet")
    document.add_paragraph(
        report["explainability"]["signed_contribution_note"]
    )

    document.add_heading(
        "Government Evidence & Standards",
        level=1,
    )
    government = report["government_evidence"]
    if not government["records"]:
        document.add_paragraph(
            "No verified government/regulatory evidence available for this prediction."
        )
    else:
        table = document.add_table(rows=1, cols=5)
        for index, value in enumerate(
            [
                "Evidence / requirement",
                "Measured value",
                "Threshold",
                "Compliance",
                "Source",
            ]
        ):
            table.rows[0].cells[index].text = value

        for item in government["records"]:
            cells = table.add_row().cells
            cells[0].text = str(item.get("evidence_or_requirement") or "N/A")
            cells[1].text = str(item.get("current_measured_value") or "N/A")
            cells[2].text = str(item.get("applicable_threshold") or "N/A")
            cells[3].text = str(item.get("compliance_status") or "N/A")
            cells[4].text = str(item.get("source_reference") or "N/A")
        document.add_paragraph(government["threshold_note"])

    document.add_heading("Recommendations", level=1)
    if not report["recommendations"]:
        document.add_paragraph(
            "No recommendation is available from the current backend recommendation logic."
        )
    for item in report["recommendations"]:
        document.add_heading(
            f"{item['priority']} · {item['recommendation']}",
            level=2,
        )
        document.add_paragraph(f"Reason: {item['reason']}")
        document.add_paragraph(
            "Supporting evidence: "
            + (" | ".join(item.get("supporting_evidence") or []) or "N/A")
        )
        document.add_paragraph(
            "Expected benefit: "
            + (
                item.get("expected_benefit")
                or "Not quantified by current recommendation logic"
            )
        )
        document.add_paragraph(
            "Suggested timeframe: "
            + (
                item.get("suggested_timeframe")
                or "Not specified by current recommendation logic"
            )
        )

    document.add_heading("Supporting Data", level=1)
    if report["supporting_data"]:
        table = document.add_table(rows=1, cols=6)
        for index, value in enumerate(
            [
                "Input",
                "Current value",
                "Expected / historical range",
                "Status",
                "Trend",
                "Last updated",
            ]
        ):
            table.rows[0].cells[index].text = value

        for item in report["supporting_data"]:
            cells = table.add_row().cells
            current = item.get("current_value")
            if current is not None and item.get("unit"):
                current = f"{current} {item['unit']}"
            cells[0].text = str(item.get("name") or "N/A")
            cells[1].text = str(current if current is not None else "N/A")
            cells[2].text = str(item.get("expected_range") or "N/A")
            cells[3].text = str(item.get("status") or "N/A")
            cells[4].text = str(item.get("trend") or "N/A")
            cells[5].text = str(item.get("last_updated") or "N/A")
    else:
        document.add_paragraph(
            "No sensor/environment supporting-data rows are available."
        )

    for item in report["supporting_data"][:8]:
        chart = _chart_png(
            item["name"],
            item.get("history", []),
            item.get("unit") or "Value",
        )
        if chart:
            document.add_heading(
                f"{item['name']} trend",
                level=2,
            )
            document.add_picture(BytesIO(chart), width=Inches(6.4))

    document.add_heading(
        "Prediction Transparency",
        level=1,
    )
    transparency = report["transparency"]
    table = document.add_table(rows=0, cols=2)
    rows = (
        ("Prediction", transparency.get("prediction") or "N/A"),
        (
            "Confidence",
            (
                f"{round(transparency['confidence'] * 100)}%"
                if transparency.get("confidence") is not None
                else "N/A"
            ),
        ),
        (
            "Model",
            f"{transparency.get('model_name') or 'N/A'} · {transparency.get('model_version') or 'N/A'}",
        ),
        ("Prediction method", transparency.get("prediction_method") or "N/A"),
        ("Prediction generated", transparency.get("prediction_generated_at") or "N/A"),
        ("Data points used", transparency.get("data_points_used") or 0),
        ("Evidence records available", transparency.get("evidence_record_count") or 0),
        ("Model stage", transparency.get("model_stage") or "N/A"),
    )
    for name, value in rows:
        cells = table.add_row().cells
        cells[0].text = str(name)
        cells[1].text = str(value)

    document.add_paragraph(
        "Traceability boundary: populated values come from the current database, canonical Twin/ML output, stored histories, and authoritative source metadata. Missing regulatory thresholds, structural failure probabilities, and signed XAI contributions are not fabricated."
    )

    for section in document.sections:
        section.top_margin = Inches(0.55)
        section.bottom_margin = Inches(0.55)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)

    for style in document.styles:
        if hasattr(style, "font"):
            style.font.name = "Arial"
            if style.font.size is None:
                style.font.size = Pt(9)

    output = BytesIO()
    document.save(output)
    return output.getvalue()
